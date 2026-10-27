
import csv
import os
import re
import sys
import tempfile
import subprocess
from datetime import datetime
from pathlib import Path
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook
import openpyxl
import pdfplumber

try:
    from docx import Document
except ImportError:
    Document = None


DAY_NAMES = {
    "monday": "Monday", "mon": "Monday",
    "tuesday": "Tuesday", "tue": "Tuesday", "tues": "Tuesday",
    "wednesday": "Wednesday", "wed": "Wednesday",
    "thursday": "Thursday", "thu": "Thursday", "thur": "Thursday", "thurs": "Thursday",
    "friday": "Friday", "fri": "Friday",
    "saturday": "Saturday", "sat": "Saturday",
    "sunday": "Sunday", "sun": "Sunday",
}

DAY_ORDER = {
    "Monday": 1, "Tuesday": 2, "Wednesday": 3, "Thursday": 4,
    "Friday": 5, "Saturday": 6, "Sunday": 7
}

ROMAN = {
    "I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6,
    "VII": 7, "VIII": 8
}


def clean(x):
    if x is None:
        return ""
    return re.sub(r"\s+", " ", str(x).replace("\xa0", " ").strip())


def clean_multiline(x):
    """Normalize text while preserving meaningful line breaks inside timetable cells."""
    if x is None:
        return ""
    lines = [re.sub(r"[ \t]+", " ", str(line).replace("\xa0", " ").strip()) for line in str(x).replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def norm(x):
    return re.sub(r"[^a-z0-9]+", "", clean(x).lower())


def normalize_roll(x):
    return re.sub(r"[\s-]+", "", clean(x).upper())


def is_real_student_roll(x):
    s = normalize_roll(x)
    return bool(re.fullmatch(r"[A-Z]{1,4}\d+[A-Z]?\d*", s))


def find_column(columns, candidates):
    mapping = {norm(c): c for c in columns}
    for c in candidates:
        if norm(c) in mapping:
            return mapping[norm(c)]
    return None


def looks_like_subject_file(path):
    return path.suffix.lower() in {".xlsx", ".xls", ".csv"}


def is_timetable_name(path):
    n = norm(path.stem)
    return "timetable" in n or "timetable" in n or "schedule" in n or "timeslot" in n


def infer_subject_name(path):
    stem = re.sub(r"\s*\(\d+\)\s*$", "", path.stem.strip())
    # Keep the user's filename as the default subject label.
    # If the filename is "Customer Relationship Management - CRM", prefer CRM.
    parts = re.split(r"\s*[-_]\s*", stem)
    if len(parts) > 1:
        last = clean(parts[-1])
        if 1 <= len(last.split()) <= 4:
            return last
    return stem


def sheet_has_student_structure(df):
    cols = list(df.columns)
    roll = find_column(cols, ["Roll No", "Roll Number", "Roll", "Registration No", "Reg No", "ID"])
    name = find_column(cols, ["Name", "Student Name", "Student"])
    return roll is not None or name is not None


def parse_subject_file(path):
    records = []
    subject = infer_subject_name(path)

    if path.suffix.lower() == ".csv":
        sheets = {"CSV": pd.read_csv(path)}
    else:
        sheets = pd.read_excel(path, sheet_name=None)

    for sheet_name, df in sheets.items():
        if df is None or df.empty:
            continue

        df = df.copy()
        df.columns = [clean(c) for c in df.columns]

        # Some subject workbooks contain a duplicate administrative EMAIL sheet.
        # It has the same students as the real enrolment sheet, but it is not a
        # separate section/course list. Skip it entirely.
        sheet_key = norm(sheet_name)
        has_email_columns = any(
            "email" in norm(c) for c in df.columns
        )
        if sheet_key in {"email", "emails"} or has_email_columns:
            continue

        roll_col = find_column(
            df.columns,
            ["Roll No", "Roll Number", "Roll", "Registration No", "Reg No", "Student ID", "ID"]
        )
        name_col = find_column(
            df.columns,
            ["Name", "Student Name", "Student", "Candidate Name"]
        )
        section_col = find_column(
            df.columns,
            ["Section", "Sec", "Batch", "Group", "Class", "Division"]
        )

        if not roll_col and not name_col:
            # Ignore non-student sheets, but allow a simple one-column student list.
            if len(df.columns) == 1:
                only_col = df.columns[0]
                values = [clean(v) for v in df[only_col].tolist() if clean(v)]
                if any(is_real_student_roll(v) for v in values):
                    roll_col = only_col
                else:
                    name_col = only_col
            else:
                continue

        multi_sheet = len(sheets) > 1
        sheet_section = "" if str(sheet_name).lower() in {"sheet1", "sheet2", "csv"} else clean(sheet_name)

        for _, row in df.iterrows():
            roll = normalize_roll(row[roll_col]) if roll_col else ""
            name = clean(row[name_col]) if name_col else ""
            section = clean(row[section_col]) if section_col else ""
            if not section and multi_sheet:
                # Treat only section-like sheet names as sections.
                # "LIST", "EMAIL", "SHEET1", etc. are administrative/data sheets.
                if re.fullmatch(
                    r"(?:[A-D]|[A-D]\d{1,2}|\d{1,2}[A-D]|PRE|POST)",
                    sheet_section or "",
                    flags=re.I,
                ):
                    section = sheet_section.upper()


            if roll or name:
                if roll and not is_real_student_roll(roll):
                    # Don't accidentally treat "Roll No" as data.
                    if norm(roll) in {"rollno", "rollnumber", "studentid", "id"}:
                        continue
                records.append({
                    "roll": roll,
                    "name": name,
                    "subject": subject,
                    "section": section,
                    "source_file": path.name,
                    "source_sheet": clean(sheet_name),
                })

    return records


def parse_all_subject_files(folder, timetable_path):
    records = []
    for path in sorted(folder.iterdir()):
        if path.is_file() and looks_like_subject_file(path) and path.resolve() != timetable_path.resolve():
            # Avoid treating likely master timetable workbooks as subject files.
            if is_timetable_name(path):
                continue
            # Common administrative files should not be treated as subjects.
            if norm(path.stem) in {
                "studentlist", "studentlists", "masterlist", "allstudents",
                "students", "timetable", "schedule"
            }:
                continue
            try:
                records.extend(parse_subject_file(path))
            except Exception:
                # Let the UI report ignored files later; one malformed file shouldn't crash everything.
                continue

    # De-duplicate.
    seen = set()
    out = []
    for r in records:
        key = (r["roll"], norm(r["name"]), norm(r["subject"]), norm(r["section"]))
        if key not in seen and (r["roll"] or r["name"]):
            seen.add(key)
            out.append(r)
    return out


def extract_time(text):
    text = clean(text).replace("–", "-").replace("—", "-")
    m = re.search(
        r"(\d{1,2}:\d{2}\s*(?:AM|PM)?)\s*-\s*(\d{1,2}:\d{2}\s*(?:AM|PM)?)",
        text,
        flags=re.I
    )
    if not m:
        return ""
    return f"{m.group(1).upper().replace(' ', '')}-{m.group(2).upper().replace(' ', '')}"


def period_label_from_header(text, index):
    text = clean(text)
    # Prefer Roman numeral when present.
    m = re.search(r"\b(I{1,3}|IV|V|VI|VII|VIII)\b", text.upper())
    if m:
        return m.group(1).upper()
    t = extract_time(text)
    if t:
        return t
    return f"Period {index}"


def weekday(text):
    k = clean(text).lower()
    return DAY_NAMES.get(k, "")


def convert_legacy_doc_to_docx(doc_path):
    # Windows + Microsoft Word: best support for old .doc files.
    if os.name == "nt":
        try:
            import win32com.client  # type: ignore
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(str(doc_path))
            tmp = Path(tempfile.gettempdir()) / (doc_path.stem + "_converted.docx")
            doc.SaveAs2(str(tmp), FileFormat=16)  # wdFormatXMLDocument
            doc.Close(False)
            word.Quit()
            return tmp
        except Exception:
            pass

    # Linux/macOS fallback if LibreOffice is installed.
    for command in ("soffice", "libreoffice"):
        try:
            out_dir = Path(tempfile.mkdtemp(prefix="tt_"))
            subprocess.run(
                [command, "--headless", "--convert-to", "docx", "--outdir", str(out_dir), str(doc_path)],
                check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True
            )
            converted = out_dir / (doc_path.stem + ".docx")
            if converted.exists():
                return converted
        except Exception:
            continue

    return None


def load_docx_table(path):
    actual = Path(path)
    temp_converted = None

    if actual.suffix.lower() == ".doc":
        converted = convert_legacy_doc_to_docx(actual)
        if converted is None:
            raise RuntimeError(
                "This is an old .doc timetable. On Windows, Microsoft Word must be installed "
                "for the program to read .doc files. Another option is to save the timetable as .docx."
            )
        actual = converted
        temp_converted = converted

    if Document is None:
        raise RuntimeError("python-docx is not installed.")

    try:
        doc = Document(str(actual))
        return doc.tables
    finally:
        if temp_converted and temp_converted.exists():
            try:
                temp_converted.unlink()
            except Exception:
                pass


def find_timetable_word_table(path):
    tables = load_docx_table(path)
    best = None
    best_score = -1

    for table in tables:
        if not table.rows:
            continue
        score = 0
        first = " ".join(clean(c.text).lower() for c in table.rows[0].cells)
        if "day" in first:
            score += 3
        if extract_time(first):
            score += 2
        if len(table.rows[0].cells) >= 4:
            score += 1
        if score > best_score:
            best_score = score
            best = table

    if best is None:
        raise RuntimeError("Could not find the timetable table in the Word file.")
    return best


def parse_matrix_timetable_from_doc(path):
    table = find_timetable_word_table(path)
    rows = table.rows
    header = [clean(c.text) for c in rows[0].cells]

    periods = []
    for i, h in enumerate(header[1:], start=1):
        periods.append({
            "index": i,
            "label": period_label_from_header(h, i),
            "time": extract_time(h),
        })

    matrix = []
    for row in rows[1:]:
        raw_cells = [str(c.text or "") for c in row.cells]
        if not raw_cells:
            continue
        day = weekday(raw_cells[0])
        if not day:
            continue
        for i in range(1, min(len(raw_cells), len(periods) + 1)):
            cell = raw_cells[i]
            if not clean(cell):
                continue
            matrix.append({
                "day": day,
                "period": periods[i - 1]["label"],
                "time": periods[i - 1]["time"],
                "cell": cell
            })
    return matrix


def find_timetable_excel_table(path):
    sheets = pd.read_excel(path, sheet_name=None)
    # First try long format.
    for sheet_name, df in sheets.items():
        if df is None or df.empty:
            continue
        df = df.copy()
        df.columns = [clean(c) for c in df.columns]
        day_col = find_column(df.columns, ["Day", "Weekday", "Week Day"])
        subject_col = find_column(df.columns, ["Subject", "Course", "Course Name", "Paper", "Class"])
        period_col = find_column(df.columns, ["Period", "Period No", "Slot", "Time", "Time Slot"])
        if day_col and subject_col and period_col:
            rows = []
            start_col = find_column(df.columns, ["Start Time", "Start", "From"])
            end_col = find_column(df.columns, ["End Time", "End", "To"])
            section_col = find_column(df.columns, ["Section", "Sec", "Batch", "Group"])
            room_col = find_column(df.columns, ["Room", "Classroom", "Venue", "Location"])
            for _, r in df.iterrows():
                day = weekday(r[day_col])
                subject = clean(r[subject_col])
                if not day or not subject:
                    continue
                rows.append({
                    "day": day,
                    "period": clean(r[period_col]),
                    "time": (
                        f"{clean(r[start_col])}-{clean(r[end_col])}"
                        if start_col and end_col else clean(r[period_col])
                    ),
                    "subject": subject,
                    "section": clean(r[section_col]) if section_col else "",
                    "room": clean(r[room_col]) if room_col else "",
                    "cell": subject,
                    "long_format": True,
                })
            if rows:
                return rows

    # Otherwise treat the sheet as a matrix like the Word timetable.
    for sheet_name, df in sheets.items():
        if df is None or df.empty or df.shape[1] < 3:
            continue
        df = df.copy()
        df.columns = [clean(c) for c in df.columns]
        header = [clean(c) for c in df.columns]
        if norm(header[0]) != "day" and "day" not in norm(header[0]):
            continue
        periods = [
            {"label": period_label_from_header(h, i), "time": extract_time(h)}
            for i, h in enumerate(header[1:], start=1)
        ]
        rows = []
        for _, r in df.iterrows():
            day = weekday(r.iloc[0])
            if not day:
                continue
            for i, p in enumerate(periods, start=1):
                cell = clean(r.iloc[i])
                if cell:
                    rows.append({
                        "day": day,
                        "period": p["label"],
                        "time": p["time"],
                        "cell": cell
                    })
        if rows:
            return rows

    raise RuntimeError("Could not detect the timetable structure in the Excel file.")



def parse_matrix_timetable_from_pdf(path):
    """Parse a text-based timetable PDF using pdfplumber table extraction."""
    candidate_tables = []

    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            settings_list = [
                {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
                {},
                {"vertical_strategy": "lines", "horizontal_strategy": "text"},
            ]
            for settings in settings_list:
                try:
                    tables = page.extract_tables(settings)
                except Exception:
                    tables = []
                for table in tables:
                    if not table:
                        continue
                    rows = [[clean_multiline(c) for c in (row or [])] for row in table]
                    if not rows:
                        continue

                    # Score tables rather than blindly selecting the widest one.
                    # The desired timetable table normally starts with DAY and has
                    # 6+ time-bearing period columns.
                    first = rows[0]
                    first_text = " ".join(first).replace("\n", " ")
                    day_first = bool(first and norm(first[0]) == "day")
                    time_cells = sum(1 for c in first if extract_time(c))
                    day_rows = sum(1 for r in rows if r and weekday(r[0]))
                    score = (
                        1000 if day_first else 0
                    ) + time_cells * 100 + min(len(first), 10) * 5 + min(day_rows, 7) * 20

                    if day_first or ("day" in norm(first_text) and extract_time(first_text)):
                        candidate_tables.append((score, page_no, rows))

    if not candidate_tables:
        raise RuntimeError(
            "Could not identify a timetable table in the PDF. "
            "Use a text-based PDF with days on the left and periods/times across the top."
        )

    _, page_no, rows = max(candidate_tables, key=lambda x: x[0])

    # The selected standard table should have the header in its first row. If not,
    # locate the best row containing DAY + time information.
    header_idx = 0
    if not (rows[0] and norm(rows[0][0]) == "day"):
        for i, row in enumerate(rows[:12]):
            row_text = " ".join(row)
            if row and norm(row[0]) == "day" and extract_time(row_text):
                header_idx = i
                break

    header = rows[header_idx]
    if len(header) < 3:
        raise RuntimeError("The PDF timetable does not contain enough period columns.")

    ncols = len(header)

    # The supplied master timetable PDF uses two header rows:
    # row 1 contains I/II/III/... and row 2 contains the actual times.
    # Combine them so the generated personal timetable gets both period and time.
    time_row = None
    if header_idx + 1 < len(rows):
        candidate = rows[header_idx + 1]
        if len(candidate) >= ncols and any(extract_time(c or "") for c in candidate):
            time_row = candidate

    periods = []
    for idx, h in enumerate(header[1:], start=1):
        h = clean(h)
        time_text = clean(time_row[idx] if time_row is not None and idx < len(time_row) else "")
        time_value = extract_time(time_text) or extract_time(h)
        period_value = re.sub(r"[^IVX]+", "", h.upper()) if re.fullmatch(r"[IVX]+", h.upper()) else h
        periods.append({
            "label": period_value or f"Period {idx}",
            "time": time_value,
        })

    # Coalesce continuation rows created by merged PDF cells.
    timetable_rows = []
    current = None
    for raw_row in rows[header_idx + 1:]:
        row = list(raw_row or []) + [""] * max(0, ncols - len(raw_row or []))
        row = [clean_multiline(c) for c in row[:ncols]]
        if not any(row):
            continue

        row_day = weekday(row[0]) if row else ""
        if row_day:
            if current is not None:
                timetable_rows.append(current)
            current = {"day": row_day, "cells": row[1:]}
        elif current is not None:
            for j in range(1, ncols):
                extra = row[j]
                if not extra:
                    continue
                if current["cells"][j - 1]:
                    current["cells"][j - 1] += "\n" + extra
                else:
                    current["cells"][j - 1] = extra

    if current is not None:
        timetable_rows.append(current)

    output = []
    for tr in timetable_rows:
        day = tr["day"]
        for j, period in enumerate(periods):
            cell = clean_multiline(tr["cells"][j] if j < len(tr["cells"]) else "")
            if not cell:
                continue
            output.append({
                "day": day,
                "period": period["label"],
                "time": period["time"],
                "cell": cell,
                "pdf_page": page_no,
            })

    if not output:
        raise RuntimeError("The PDF timetable table was found, but no class cells were extracted.")

    return output



def parse_timetable(path):
    path = Path(path)
    if not path.exists():
        raise RuntimeError(
            f"Timetable file was not found at:\n{path}\n\n"
            "Check that the file still exists and that the folder was selected correctly."
        )
    ext = path.suffix.lower()
    if ext in {".doc", ".docx"}:
        return parse_matrix_timetable_from_doc(path)
    if ext == ".pdf":
        return parse_matrix_timetable_from_pdf(path)
    if ext in {".xlsx", ".xls"}:
        return find_timetable_excel_table(path)
    if ext == ".csv":
        tmp = pd.read_csv(path)
        # Reuse Excel logic by writing a temporary xlsx.
        temp = Path(tempfile.gettempdir()) / "_tt_temp.xlsx"
        tmp.to_excel(temp, index=False)
        try:
            return find_timetable_excel_table(temp)
        finally:
            try:
                temp.unlink()
            except Exception:
                pass
    raise RuntimeError("Unsupported timetable format.")


def subject_matches_cell(subject, section, cell):
    """Match a subject entry without accidentally matching another section."""
    text = clean(cell)
    s = clean(subject)
    sec = clean(section)
    if not s:
        return False

    subj_re = re.escape(s)

    # The subject itself must occur as a token.
    if not re.search(rf"(?<![A-Za-z0-9]){subj_re}(?![A-Za-z0-9])", text, flags=re.I):
        return False

    if not sec:
        return True

    sec_re = re.escape(sec)

    # Style 1: BrM(D), BrM (D)
    if re.search(
        rf"(?<![A-Za-z0-9]){subj_re}\s*\(\s*{sec_re}\s*\)",
        text, flags=re.I
    ):
        return True

    # Style 2: SO (AP) B L-3 / SO B L-3
    # Permit one parenthesized faculty token between subject and section,
    # but require the section to be the next meaningful token. This prevents
    # BrM(B) from incorrectly matching BrM(C).
    if re.search(
        rf"(?<![A-Za-z0-9]){subj_re}\s+(?:\([^)]*\)\s+)?{sec_re}(?![A-Za-z0-9])",
        text, flags=re.I
    ):
        return True

    return False

def matching_line(cell, subject, section):
    # Word tables often place multiple classes in one cell using line breaks.
    # Keep the line boundaries so the room belongs to the matched subject.
    raw = str(cell or "").replace("\r", "\n")
    lines = [clean(x) for x in re.split(r"\n|\|", raw) if clean(x)]

    for line in lines:
        if subject_matches_cell(subject, section, line):
            return line

    if subject_matches_cell(subject, section, clean(raw)):
        return clean(raw)

    return ""


def extract_room(cell):
    text = str(cell or "")
    # Locations in the timetable include G-12, G-14, L-B and LB.
    m = re.search(
        r"\b(?:G|L)\s*-\s*\d+\b|\bL\s*-\s*B\b|\bLB\b",
        text,
        flags=re.I
    )
    return clean(m.group(0)) if m else ""


def extract_faculty(cell, subject="", section=""):
    """
    Extract faculty initials from a timetable line such as:
        CRM (ANK/NB) G-12
        BrM (D) (AK)
        BrM(D) (AK) G-12
        SO (AP) B L-3

    When a section is present in parentheses, ignore that parenthesis and
    return the faculty initials.
    """
    raw = str(cell or "")
    candidates = re.findall(
        r"\(\s*([A-Za-z]+(?:\s*/\s*[A-Za-z]+)*)\s*\)",
        raw
    )
    candidates = [clean(x).replace(" ", "") for x in candidates]

    sec = clean(section).upper()
    for candidate in candidates:
        if sec and candidate.upper() == sec:
            continue

        # Faculty initials in the supplied timetable are alphabetic tokens.
        # Avoid treating words such as PRE/POST as faculty.
        if candidate.upper() in {"PRE", "POST"}:
            continue
        return candidate

    return ""



def build_student_timetable(records, timetable, roll):
    roll = normalize_roll(roll)
    student_rows = [r for r in records if normalize_roll(r["roll"]) == roll]

    if not student_rows:
        return None, [], ["Roll number not found in the subject files."]

    name = next((r["name"] for r in student_rows if r["name"]), "")

    schedule = []
    missing = []

    # Preserve the subject/section combination from the enrollment files.
    seen_courses = set()
    for r in student_rows:
        course_key = (norm(r["subject"]), norm(r["section"]))
        if course_key in seen_courses:
            continue
        seen_courses.add(course_key)

        hits = []
        for slot in timetable:
            if "long_format" in slot and slot["long_format"]:
                if norm(slot.get("subject")) != norm(r["subject"]):
                    # Allow filename code to appear in long-format subject.
                    if norm(r["subject"]) not in norm(slot.get("subject")):
                        continue
                if r["section"] and clean(slot.get("section")) and norm(slot.get("section")) != norm(r["section"]):
                    continue
                if r["section"] and not clean(slot.get("section")):
                    # Long format without section cannot safely decide.
                    continue
                hit = dict(slot)
                hit["subject"] = r["subject"]
                hit["section"] = r["section"]
                hit["room"] = slot.get("room", "")
                hits.append(hit)
            else:
                line = matching_line(slot["cell"], r["subject"], r["section"])
                if line:
                    hit = dict(slot)
                    hit["subject"] = r["subject"]
                    hit["section"] = r["section"]
                    hit["room"] = extract_room(line)
                    hit["faculty"] = extract_faculty(
                        line, subject=r["subject"], section=r["section"]
                    )
                    hits.append(hit)

        if not hits:
            missing.append(
                f'{r["subject"]}' + (f' ({r["section"]})' if r["section"] else "")
            )

        schedule.extend(hits)

    # Remove duplicate matches and sort.
    unique = []
    seen = set()
    for r in schedule:
        key = (
            r["day"], r["period"], norm(r["subject"]), norm(r.get("section", "")),
            norm(r.get("room", ""))
        )
        if key not in seen:
            seen.add(key)
            unique.append(r)

    unique.sort(key=lambda x: (
        DAY_ORDER.get(x["day"], 99),
        str(x["period"]),
        norm(x["subject"])
    ))

    return {
        "roll": roll,
        "name": name,
        "courses": student_rows,
        "schedule": unique,
        "missing": sorted(set(missing), key=str.lower),
    }, student_rows, []


def make_text(student):
    lines = [
        f'{student["name"]} — Roll No: {student["roll"]}',
        "",
        "PERSONAL TIMETABLE",
        "",
    ]

    grouped = defaultdict(list)
    for x in student["schedule"]:
        grouped[x["day"]].append(x)

    for day_num in sorted(DAY_ORDER.values()):
        day_name = next(
            (name for name, number in DAY_ORDER.items() if number == day_num),
            None
        )
        if not day_name or day_name not in grouped:
            continue

        lines.append(day_name.upper())
        lines.append(
            "Time                  Subject              Section   Faculty     Location"
        )
        lines.append(
            "--------------------  -------------------  -------   ----------  --------"
        )

        for x in grouped[day_name]:
            time = x.get("time", "") or x.get("period", "")
            subject = clean(x.get("subject", ""))
            section = clean(x.get("section", "")) or "—"
            faculty = clean(x.get("faculty", "")) or "—"
            location = clean(x.get("room", "")) or "—"

            lines.append(
                f'{time:<20}  {subject:<19}  {section:<7}   '
                f'{faculty:<10}  {location:<8}'
            )
        lines.append("")

    if student["missing"]:
        lines.append("SUBJECTS WITH NO MATCHING TIMETABLE ENTRY")
        for s in student["missing"]:
            lines.append(f"- {s}")
        lines.append("")

    if not student["schedule"]:
        lines.append("No timetable entries were found for this student.")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"



def make_master_grid_text(student):
    master_periods = [
        ("I", "8:30-10:00 AM"),
        ("II", "10:15-11:45 AM"),
        ("III", "12:00-1:30 PM"),
        ("IV", "2:15-3:45 PM"),
        ("V", "4:00-5:30 PM"),
        ("VI", "5:45-7:15 PM"),
    ]

    lookup = defaultdict(list)
    for x in student["schedule"]:
        period = clean(x.get("period", "")).upper()
        if period not in {p for p, _ in master_periods}:
            tkey = re.sub(r"[^0-9apm]", "", clean(x.get("time", "")).lower())
            for p, t in master_periods:
                if re.sub(r"[^0-9apm]", "", t.lower()) == tkey:
                    period = p
                    break

        details = [clean(x["subject"])]
        meta = []
        if clean(x.get("section", "")):
            meta.append(f"Sec {clean(x['section'])}")
        if clean(x.get("faculty", "")):
            meta.append(clean(x["faculty"]))
        if clean(x.get("room", "")):
            meta.append(clean(x["room"]))
        if meta:
            details.append(" | ".join(meta))

        lookup[(x["day"], period)].append(" — ".join(details))

    lines = [
        f"PERSONAL TIME TABLE — {student['name']}",
        f"Roll No: {student['roll']}",
        "",
        "DAY | " + " | ".join(f"{p}\n{t}" for p, t in master_periods),
    ]

    for day in ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]:
        cells = [
            " / ".join(lookup.get((day, p), []))
            for p, _ in master_periods
        ]
        lines.append(day + " | " + " | ".join(cells))

    return "\n".join(lines) + "\n\n"


def save_outputs(student, folder):
    out_dir = folder / "Student Timetables"
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_roll = re.sub(r"[^A-Za-z0-9_-]+", "_", student["roll"])
    safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "_", student["name"]).strip() or "Student"

    # The supplied master timetable has a fixed six-period structure.
    # Keep these columns exactly, even when a student has no class in a period.
    master_periods = [
        ("I", "8:30-10:00 AM"),
        ("II", "10:15-11:45 AM"),
        ("III", "12:00-1:30 PM"),
        ("IV", "2:15-3:45 PM"),
        ("V", "4:00-5:30 PM"),
        ("VI", "5:45-7:15 PM"),
    ]

    def time_key(value):
        return re.sub(r"[^0-9apm]", "", clean(value).lower().replace("–", "-").replace("—", "-"))

    master_by_time = {time_key(t): p for p, t in master_periods}
    master_by_period = {p.upper(): p for p, _ in master_periods}

    rows = []
    for x in student["schedule"]:
        source_time = clean(x.get("time", ""))
        source_period = clean(x.get("period", "")).upper()

        # Prefer time because the PDF header has a separate period row and time row.
        canonical_period = master_by_time.get(time_key(source_time), "")
        if not canonical_period:
            canonical_period = master_by_period.get(source_period, "")

        # If an unusual source format has a time written inside the period field,
        # match that too.
        if not canonical_period:
            canonical_period = master_by_time.get(time_key(source_period), "")

        # Fall back gracefully; regular master timetable inputs will map to I-VI.
        if not canonical_period:
            canonical_period = source_period or "UNKNOWN"

        rows.append({
            "Day": x["day"],
            "Time": source_time,
            "Period": canonical_period,
            "Subject": x["subject"],
            "Section": x.get("section", ""),
            "Faculty": x.get("faculty", ""),
            "Location": x.get("room", ""),
        })

    df = pd.DataFrame(
        rows,
        columns=["Day", "Time", "Period", "Subject", "Section", "Faculty", "Location"]
    )

    xlsx_path = out_dir / f"{safe_roll}_{safe_name}_Timetable.xlsx"
    txt_path = out_dir / f"{safe_roll}_{safe_name}_Timetable.txt"

    # Build class-cell text.
    cell_lookup = defaultdict(list)
    for _, r in df.iterrows():
        details = [
            clean(r["Subject"]),
            " | ".join(
                p for p in [
                    f"Section: {clean(r['Section'])}" if clean(r["Section"]) else "",
                    f"Faculty: {clean(r['Faculty'])}" if clean(r["Faculty"]) else "",
                    f"Location: {clean(r['Location'])}" if clean(r["Location"]) else "",
                ]
                if p
            ),
        ]
        cell_lookup[(clean(r["Day"]), clean(r["Period"]))].append(
            "\n".join(p for p in details if p)
        )

    # Full master timetable day structure.
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"

    # Match the master timetable's hierarchy.
    ws.merge_cells("A1:G1")
    ws["A1"] = "MANAGEMENT DEVELOPMENT INSTITUTE GURGAON"
    ws.merge_cells("A2:G2")
    ws["A2"] = f"PERSONAL TIME TABLE — {student['name']}"
    ws.merge_cells("A3:G3")
    ws["A3"] = f"Roll No: {student['roll']}"

    ws["A5"] = "DAY"
    for c, (period, time) in enumerate(master_periods, start=2):
        ws.cell(5, c, period)
        ws.cell(6, c, time)
    ws["A6"] = ""

    # Insert exact five weekday rows.
    for row_idx, day in enumerate(days, start=7):
        ws.cell(row_idx, 1, day)
        for col_idx, (period, _) in enumerate(master_periods, start=2):
            entries = cell_lookup.get((day, period), [])
            ws.cell(row_idx, col_idx, "\n---\n".join(entries))

    # Details sheet.
    details_ws = wb.create_sheet("Details")
    headers = ["Day", "Time", "Period", "Subject", "Section", "Faculty", "Location"]
    for c, h in enumerate(headers, start=1):
        details_ws.cell(1, c, h)
    for r_idx, (_, r) in enumerate(df.iterrows(), start=2):
        for c_idx, h in enumerate(headers, start=1):
            details_ws.cell(r_idx, c_idx, r[h])

    # Formatting to resemble the master timetable grid.
    thin = openpyxl.styles.Side(style="thin")
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=5, max_row=11, min_col=1, max_col=7):
        for cell in row:
            cell.border = border
            cell.alignment = openpyxl.styles.Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=7):
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    ws["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    ws["A2"].font = openpyxl.styles.Font(bold=True, size=13)
    ws["A3"].font = openpyxl.styles.Font(bold=True, size=11)

    ws.column_dimensions["A"].width = 17
    for col in range(2, 8):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 23

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[5].height = 23
    ws.row_dimensions[6].height = 32
    for r in range(7, 12):
        ws.row_dimensions[r].height = 82

    ws.freeze_panes = "B7"
    ws.sheet_view.showGridLines = False
    ws.print_area = "A1:G11"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    details_ws.freeze_panes = "A2"
    details_ws.auto_filter.ref = details_ws.dimensions
    details_ws.sheet_view.showGridLines = False
    for c in range(1, 8):
        details_ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 22

    wb.save(xlsx_path)

    # Copyable text uses the same master-grid representation.
    master_text = make_master_grid_text(student)

    txt_path.write_text(master_text, encoding="utf-8")
    return xlsx_path, txt_path





import streamlit as st
import html
import tempfile
import shutil
from pathlib import Path
from collections import defaultdict

st.set_page_config(page_title="Student Timetable Generator", page_icon="📅", layout="wide")

DATA_DIR = Path(__file__).parent / "data"
MASTER_PERIODS = [
    ("I","8:30–10:00 AM"),
    ("II","10:15–11:45 AM"),
    ("III","12:00–1:30 PM"),
    ("IV","2:15–3:45 PM"),
    ("V","4:00–5:30 PM"),
    ("VI","5:45–7:15 PM"),
]

def find_master_timetable():
    """Find the class timetable without accidentally choosing an exam date sheet."""
    pdfs = sorted(DATA_DIR.glob("*.pdf"))
    class_pdfs = [
        p for p in pdfs
        if not any(term in norm(p.stem) for term in ("exam", "datesheet", "examination"))
    ]
    named = [
        p for p in class_pdfs
        if re.search(r"time[\s_-]*table|timetable|schedule", p.stem, flags=re.I)
    ]
    return (named or class_pdfs or [None])[0]


def find_exam_datesheet():
    """Find the exam date-sheet PDF bundled in the repository's data folder."""
    pdfs = sorted(DATA_DIR.glob("*.pdf"))
    matches = [
        p for p in pdfs
        if ("exam" in norm(p.stem) and ("date" in norm(p.stem) or "sheet" in norm(p.stem)))
        or "datesheet" in norm(p.stem)
    ]
    return matches[0] if matches else None


EXAM_SUBJECT_NAMES = {
    "brm": "Brand Management",
    "mvw": "Marketing for Virtual World",
    "crm": "Customer Relationship Management",
    "bs": "Behavioral Strategy",
    "og": "Organisational Governance",
    "in": "Innovation Strategy",
    "lmgb": "Locational AI, Mapping and Geospatial Business",
    "so": "Sustainable Operations",
    "sdt": "Systems and Design Thinking",
    "bfbv": "Behavioural Finance and Business Valuation",
    "pfp": "Personal Financial Planning",
    "psf": "Project & Structured Finance",
    "gees": "Global Economic Environment & Strategy",
    "im": "International Marketing",
    "mdt": "Managing Digital Transformation",
    "nlpb": "Natural Language Processing for Business",
    "rmbfi": "Risk Management in Banks and Financial Institutions",
    "sm": "Sports Marketing",
    "sdw": "Strategies in a Deglobalized World",
}

EXAM_DISPLAY_CODES = {
    "brm": "BrM", "mvw": "MVW", "crm": "CRM", "bs": "BS", "og": "OG",
    "in": "IN", "lmgb": "LMGB", "so": "SO", "sdt": "SDT", "bfbv": "BFBV",
    "pfp": "PFP", "psf": "PSF", "gees": "GEES", "im": "IM", "mdt": "MDT",
    "nlpb": "NLPB", "rmbfi": "RMBFI", "sm": "SM", "sdw": "SDW",
}

EXAM_NAME_ALIASES = {
    norm(name): code for code, name in EXAM_SUBJECT_NAMES.items()
}
EXAM_NAME_ALIASES.update({
    norm("Behavioral Finance and Business Valuation"): "bfbv",
    norm("Organizational Governance"): "og",
    norm("Global Economic Environment and Strategy"): "gees",
    norm("Strategies in a Deglobalised World"): "sdw",
    norm("Locational AI Mapping and Geospatial Business"): "lmgb",
})


def canonical_exam_code(subject):
    key = norm(subject)
    if key in EXAM_SUBJECT_NAMES:
        return key
    return EXAM_NAME_ALIASES.get(key, "")


def parse_exam_datesheet_pdf(path):
    """Read the Day/Date/Time table from the institute's exam date-sheet PDF."""
    exam_rows = []
    header_terms = ("10.15 am", "2.15 pm", "6.00 pm")
    day_set = set(DAY_ORDER)
    date_pattern = re.compile(r"([A-Za-z]+\s+\d{1,2},\s*\d{4})")

    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            for table in page.find_tables():
                rows = table.extract()
                if not rows:
                    continue

                header_idx = None
                for i, row in enumerate(rows[:5]):
                    header_text = " ".join(clean(c).lower() for c in row if c)
                    header_norm = header_text.replace(":", ".")
                    if "day" in header_text and "date" in header_text and all(t in header_norm for t in header_terms):
                        header_idx = i
                        break
                if header_idx is None:
                    continue

                header = rows[header_idx]
                header_cells = table.rows[header_idx].cells
                time_centers = []
                for idx, value in enumerate(header):
                    label = clean(value).lower().replace(":", ".")
                    if "10.15 am" in label:
                        time_label = "10:15 AM"
                    elif "2.15 pm" in label:
                        time_label = "2:15 PM"
                    elif "6.00 pm" in label:
                        time_label = "6:00 PM"
                    else:
                        continue
                    bbox = header_cells[idx] if idx < len(header_cells) else None
                    if bbox:
                        time_centers.append((time_label, (bbox[0] + bbox[2]) / 2))

                if len(time_centers) < 3:
                    continue

                for row_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 1):
                    day = next((clean(c) for c in row if clean(c).title() in day_set), "")
                    date_match = next((date_pattern.search(clean(c)) for c in row if date_pattern.search(clean(c))), None)
                    if not day or not date_match:
                        continue
                    date_text = date_match.group(1)
                    try:
                        exam_date = datetime.strptime(date_text, "%B %d, %Y").date()
                    except ValueError:
                        continue

                    data_cells = table.rows[row_idx].cells
                    for idx, value in enumerate(row):
                        cell_text = clean(value)
                        if not cell_text or cell_text.title() in day_set or date_pattern.search(cell_text):
                            continue
                        # The exam grid uses merged cells. Use each cell's horizontal
                        # centre to identify whether it belongs to the 10:15, 2:15,
                        # or 6:00 PM session, preserving empty sessions on weekends.
                        bbox = data_cells[idx] if idx < len(data_cells) else None
                        if not bbox:
                            continue
                        center = (bbox[0] + bbox[2]) / 2
                        nearest_time, nearest_center = min(time_centers, key=lambda x: abs(x[1] - center))
                        distance = abs(nearest_center - center)
                        if distance > 30:
                            continue
                        time_label = nearest_time

                        codes = [clean(part) for part in re.split(r"[/,;|]+", cell_text) if clean(part)]
                        for code in codes:
                            code_key = canonical_exam_code(code)
                            if not code_key:
                                continue
                            exam_rows.append({
                                "day": exam_date.strftime("%A"),
                                "date_obj": exam_date,
                                "date": exam_date.strftime("%B %d, %Y"),
                                "time": time_label,
                                "code": EXAM_DISPLAY_CODES.get(code_key, code.upper()),
                                "code_key": code_key,
                                "subject": EXAM_SUBJECT_NAMES[code_key],
                            })

                # The date-sheet table has been parsed; don't inspect the faculty table.
                if exam_rows:
                    break
            if exam_rows:
                break

    unique = []
    seen = set()
    for row in exam_rows:
        key = (row["date"], row["time"], row["code_key"])
        if key not in seen:
            seen.add(key)
            unique.append(row)

    if not unique:
        raise RuntimeError("The exam date-sheet PDF was found, but its Day/Date/Time table could not be read.")

    time_order = {"10:15 AM": 1, "2:15 PM": 2, "6:00 PM": 3}
    unique.sort(key=lambda x: (x["date_obj"], time_order.get(x["time"], 99), x["code"]))
    return unique


def build_student_exam_timetable(student, all_exams):
    """Filter the exam date sheet to subjects found for this student's roll number."""
    enrolled = {}
    for record in student.get("courses", []):
        code = canonical_exam_code(record.get("subject", ""))
        if code:
            enrolled[code] = EXAM_DISPLAY_CODES.get(code, record.get("subject", ""))

    exams = [row for row in all_exams if row["code_key"] in enrolled]
    found_codes = {row["code_key"] for row in exams}
    unmapped = [
        clean(record.get("subject", ""))
        for record in student.get("courses", [])
        if not canonical_exam_code(record.get("subject", ""))
    ]
    missing = [EXAM_DISPLAY_CODES.get(code, code.upper()) for code in enrolled if code not in found_codes]
    return exams, sorted(set(missing)), sorted(set(unmapped))


def render_exam_table(exams, student):
    headers = ["DAY", "DATE", "TIME", "SUBJECT", "COURSE NAME"]
    header_html = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body_rows = []
    for item in exams:
        values = [item["day"], item["date"], item["time"], item["code"], item["subject"]]
        body_rows.append("<tr>" + "".join(f"<td>{html.escape(str(v))}</td>" for v in values) + "</tr>")
    if not body_rows:
        body_rows = ['<tr><td colspan="5" class="empty">No matching exams were found for the enrolled subjects.</td></tr>']
    return f'''
<style>
.exam-wrap{{width:100%;overflow-x:auto;color:#f0f2f6}}
.exam-table{{width:100%;border-collapse:collapse;table-layout:auto;background:#0e1117;color:#f0f2f6}}
.exam-table th,.exam-table td{{border:1px solid #3b414b;padding:12px 14px;text-align:left;vertical-align:middle}}
.exam-table th{{background:#1b222c;color:#f0f2f6;font-size:12px;letter-spacing:.04em}}
.exam-table td{{background:#0e1117;color:#f0f2f6;font-size:14px}}
.exam-table tr:nth-child(even) td{{background:#11161d}}
.exam-table td:nth-child(4){{font-weight:700;color:#ffffff;white-space:nowrap}}
.exam-table .empty{{text-align:center;color:#c9d1d9;padding:24px}}
@media (max-width:700px){{.exam-table th,.exam-table td{{padding:8px;font-size:12px}}}}
</style>
<div style="text-align:center;margin:6px 0 14px 0">
<div style="font-size:21px;font-weight:700">PERSONAL EXAM TIMETABLE</div>
<div style="font-size:14px;margin-top:4px">{html.escape(student["name"])} · Roll No: {html.escape(student["roll"])}</div>
<div style="font-size:12px;color:#c9d1d9;margin-top:4px">Mid-Term Examination · Dates and sessions from the official date sheet</div>
</div>
<div class="exam-wrap"><table class="exam-table"><thead><tr>{header_html}</tr></thead><tbody>{''.join(body_rows)}</tbody></table></div>
'''

def canonical_period(period,time):
    p=clean(period).upper()
    valid={x[0] for x in MASTER_PERIODS}
    if p in valid: return p
    key=re.sub(r"[^0-9apm]","",clean(time).lower())
    for mp,mt in MASTER_PERIODS:
        if key and key==re.sub(r"[^0-9apm]","",mt.lower()):
            return mp
    return p

def render_master_grid(student):
    lookup=defaultdict(list)
    for x in student["schedule"]:
        p=canonical_period(x.get("period",""),x.get("time",""))
        parts=[clean(x.get("subject",""))]
        if clean(x.get("section","")): parts.append(f"Section: {clean(x['section'])}")
        if clean(x.get("faculty","")): parts.append(f"Faculty: {clean(x['faculty'])}")
        if clean(x.get("room","")): parts.append(f"Location: {clean(x['room'])}")
        lookup[(x["day"],p)].append(parts)

    rows=[
        '<tr><th rowspan="2" class="day">DAY</th>'+''.join(
            f'<th class="period">{html.escape(p)}</th>' for p,_ in MASTER_PERIODS
        )+'</tr>',
        '<tr>'+''.join(
            f'<th class="time">{html.escape(t)}</th>' for _,t in MASTER_PERIODS
        )+'</tr>'
    ]
    for day in ["Monday","Tuesday","Wednesday","Thursday","Friday"]:
        cells=[f'<td class="day">{day}</td>']
        for p,_ in MASTER_PERIODS:
            entries=lookup.get((day,p),[])
            if not entries:
                cells.append('<td></td>'); continue
            blocks=[]
            for parts in entries:
                blocks.append(
                    '<div class="subject">'+html.escape(parts[0])+'</div>'+
                    ''.join('<div class="detail">'+html.escape(v)+'</div>' for v in parts[1:])
                )
            cells.append('<td>'+'<hr class="sep">'.join(blocks)+'</td>')
        rows.append('<tr>'+''.join(cells)+'</tr>')
    return """
<style>
.tt{width:100%;border-collapse:collapse;table-layout:fixed;background:#0e1117;color:#f0f2f6}
.tt th,.tt td{border:1px solid #3b414b;text-align:center;vertical-align:middle;padding:7px}
.tt th{color:#f0f2f6}
.tt .day{width:11%;background:#161b22;color:#f0f2f6;font-weight:700}
.tt .period{background:#1b222c;color:#f0f2f6;font-size:14px}
.tt .time{background:#11161d;color:#c9d1d9;font-size:12px}
.tt tr td:not(.day){background:#0e1117;color:#f0f2f6}
.tt td{height:92px;font-size:11px}
.subject{font-size:14px;font-weight:700;color:#ffffff;margin-bottom:4px}
.detail{font-size:11px;line-height:1.3;color:#e6edf3}
.sep{border:0;border-top:1px solid #30363d;margin:5px 0}
</style>
""" + f"""
<div style="text-align:center;margin-bottom:8px">
<div style="font-size:20px;font-weight:700">MANAGEMENT DEVELOPMENT INSTITUTE GURGAON</div>
<div style="font-size:16px;font-weight:600">PERSONAL TIME TABLE — {html.escape(student["name"])}</div>
<div style="font-size:13px">Roll No: {html.escape(student["roll"])}</div>
</div>
<table class="tt"><tbody>{''.join(rows)}</tbody></table>
"""

def _student_slot_lookup(student):
    """Build a day/period lookup for a student's parsed class schedule."""
    lookup = defaultdict(list)
    for item in student.get("schedule", []):
        period = canonical_period(item.get("period", ""), item.get("time", ""))
        lookup[(clean(item.get("day", "")), period)].append(item)
    return lookup


def render_comparison_grid(students):
    """Render a dark, horizontally scrollable comparison grid for up to six students."""
    slot_lookups = [_student_slot_lookup(person) for person in students]

    # An identical subject/section in the same slot is marked as a shared class.
    shared_people = defaultdict(set)
    for person_idx, lookup in enumerate(slot_lookups):
        for (day, period), items in lookup.items():
            for item in items:
                course_key = (norm(item.get("subject", "")), norm(item.get("section", "")))
                if course_key[0]:
                    shared_people[(day, period, course_key)].add(person_idx)

    head_cells = [
        '<th class="axis">DAY</th>',
        '<th class="axis">PERIOD</th>',
        '<th class="axis time-axis">TIME</th>',
    ]
    for idx, person in enumerate(students):
        name = clean(person.get("name", "")) or ("You" if idx == 0 else f"Friend {idx}")
        who = "You" if idx == 0 else f"Friend {idx}"
        head_cells.append(
            '<th class="person-head">' + html.escape(who) +
            '<small>' + html.escape(name + " · " + person["roll"]) + '</small></th>'
        )

    rows = []
    weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    for day in weekdays:
        for period, time_label in MASTER_PERIODS:
            cells = [
                '<td class="day-cell">' + html.escape(day) + '</td>',
                '<td class="period-cell">' + html.escape(period) + '</td>',
                '<td class="time-cell">' + html.escape(time_label) + '</td>',
            ]
            for person_idx, lookup in enumerate(slot_lookups):
                entries = lookup.get((day, period), [])
                if not entries:
                    cells.append('<td class="free-cell"><span>Free</span></td>')
                    continue
                blocks = []
                for item in entries:
                    subject = clean(item.get("subject", "")) or "Class"
                    section = clean(item.get("section", ""))
                    room = clean(item.get("room", ""))
                    faculty = clean(item.get("faculty", ""))
                    course_key = (norm(subject), norm(section))
                    is_shared = len(shared_people.get((day, period, course_key), set())) > 1
                    details = []
                    if section:
                        details.append("Section " + section)
                    if room:
                        details.append(room)
                    if faculty:
                        details.append(faculty)
                    blocks.append(
                        '<div class="course-block' + (' shared' if is_shared else '') + '">' +
                        '<div class="course-title">' + html.escape(subject) + '</div>' +
                        ('<div class="course-details">' + html.escape(" · ".join(details)) + '</div>' if details else '') +
                        ('<span class="shared-tag">SAME CLASS</span>' if is_shared else '') +
                        '</div>'
                    )
                cells.append('<td class="busy-cell">' + ''.join(blocks) + '</td>')
            rows.append('<tr>' + ''.join(cells) + '</tr>')

    return """
<style>
.compare-wrap{width:100%;overflow-x:auto;color:#f0f2f6;border:1px solid #30363d;border-radius:8px}
.compare-table{width:100%;min-width:980px;border-collapse:collapse;table-layout:auto;background:#0e1117;color:#f0f2f6}
.compare-table th,.compare-table td{border:1px solid #30363d;padding:8px 9px;text-align:center;vertical-align:middle}
.compare-table th{background:#1b222c;color:#f0f2f6;font-size:12px;white-space:nowrap}
.compare-table th small{display:block;font-size:10px;font-weight:400;color:#c9d1d9;margin-top:4px;white-space:normal}
.compare-table .axis{position:sticky;top:0;z-index:1}
.compare-table .day-cell{background:#161b22;font-weight:700;white-space:nowrap}
.compare-table .period-cell{background:#161b22;font-weight:700}
.compare-table .time-cell{background:#11161d;color:#c9d1d9;font-size:11px;min-width:112px;white-space:nowrap}
.compare-table td{min-width:130px;font-size:12px}
.compare-table .free-cell{background:#0e1117;color:#7d8590;font-size:11px}
.compare-table .busy-cell{background:#0e1117}
.course-block{position:relative;margin:2px 0;padding:6px 5px;border:1px solid #30363d;border-radius:6px;background:#11161d}
.course-block.shared{border-color:#2f855a;background:#10251d}
.course-title{font-size:13px;font-weight:700;color:#f0f2f6}
.course-details{font-size:10px;line-height:1.35;color:#c9d1d9;margin-top:3px}
.shared-tag{display:inline-block;margin-top:4px;border-radius:3px;background:#22543d;color:#d1fae5;font-size:9px;font-weight:700;padding:2px 4px}
@media(max-width:700px){.compare-table th,.compare-table td{padding:6px;font-size:11px}.course-title{font-size:12px}}
</style>
<div class="compare-wrap"><table class="compare-table"><thead><tr>""" + ''.join(head_cells) + """</tr></thead><tbody>""" + ''.join(rows) + """</tbody></table></div>
<p style="color:#9da7b3;font-size:11px;margin-top:7px">Cells marked SAME CLASS indicate students with the same subject and section at the same time. Scroll horizontally to see all students.</p>
"""


def comparison_summary(students):
    """Return identical classes and slots when every compared student is free."""
    slot_lookups = [_student_slot_lookup(person) for person in students]
    shared_people = defaultdict(set)
    for person_idx, lookup in enumerate(slot_lookups):
        for (day, period), items in lookup.items():
            for item in items:
                key = (day, period, norm(item.get("subject", "")), norm(item.get("section", "")))
                if key[2]:
                    shared_people[key].add(person_idx)

    shared_rows = []
    period_times = dict(MASTER_PERIODS)
    for (day, period, subject_key, section_key), people_idxs in shared_people.items():
        if len(people_idxs) < 2:
            continue
        subject = next(
            (clean(item.get("subject", "")) for idx in people_idxs
             for item in slot_lookups[idx].get((day, period), [])
             if norm(item.get("subject", "")) == subject_key and norm(item.get("section", "")) == section_key),
            subject_key.upper()
        )
        section = next(
            (clean(item.get("section", "")) for idx in people_idxs
             for item in slot_lookups[idx].get((day, period), [])
             if norm(item.get("subject", "")) == subject_key and norm(item.get("section", "")) == section_key),
            ""
        )
        participants = []
        for idx in sorted(people_idxs):
            person_name = clean(students[idx].get("name", "")) or ("You" if idx == 0 else f"Friend {idx}")
            participants.append(f'{person_name} ({students[idx]["roll"]})')
        shared_rows.append({
            "Day": day,
            "Period": period,
            "Time": period_times.get(period, ""),
            "Subject": subject,
            "Section": section or "—",
            "Shared by": ", ".join(participants),
        })

    day_order = {day: i for i, day in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])}
    period_order = {period: i for i, (period, _) in enumerate(MASTER_PERIODS)}
    shared_rows.sort(key=lambda row: (day_order.get(row["Day"], 99), period_order.get(row["Period"], 99), row["Subject"]))

    common_free = []
    for day in ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]:
        for period, time_label in MASTER_PERIODS:
            if all(not lookup.get((day, period), []) for lookup in slot_lookups):
                common_free.append({"Day": day, "Period": period, "Time": time_label})
    return shared_rows, common_free


def build_exam_comparison(students, all_exams):
    """Create a shared session-by-session exam view for all compared students."""
    if not all_exams:
        return pd.DataFrame()

    session_key = lambda item: (item.get("date"), item.get("time"))
    sessions = {}
    time_order = {"10:15 AM": 1, "2:15 PM": 2, "6:00 PM": 3}
    for exam in all_exams:
        key = session_key(exam)
        sessions[key] = {
            "Day": exam["day"], "Date": exam["date"], "_date_obj": exam["date_obj"], "Time": exam["time"]
        }

    column_labels = []
    per_student_exams = []
    for idx, person in enumerate(students):
        name = clean(person.get("name", "")) or ("You" if idx == 0 else f"Friend {idx}")
        label = ("You" if idx == 0 else f"Friend {idx}") + f" · {name} ({person['roll']})"
        column_labels.append(label)
        exam_rows, _, _ = build_student_exam_timetable(person, all_exams)
        student_sessions = defaultdict(list)
        for row in exam_rows:
            student_sessions[session_key(row)].append(f"{row['code']} — {row['subject']}")
        per_student_exams.append(student_sessions)

    result = []
    for key, session in sorted(sessions.items(), key=lambda pair: (pair[1]["_date_obj"], time_order.get(pair[1]["Time"], 99))):
        row = {"Day": session["Day"], "Date": session["Date"], "Time": session["Time"]}
        for label, student_sessions in zip(column_labels, per_student_exams):
            row[label] = "\n".join(student_sessions.get(key, [])) or "No exam"
        result.append(row)
    return pd.DataFrame(result)


def make_downloads(student):
    tmp=Path(tempfile.mkdtemp(prefix="ttweb_"))
    try:
        xlsx,txt=save_outputs(student,tmp)
        return xlsx.read_bytes(),txt.read_bytes()
    finally:
        shutil.rmtree(tmp,ignore_errors=True)

tt_path = find_master_timetable()
if not tt_path:
    st.error("The administrator has not installed a class timetable.")
    st.stop()

st.markdown(
    "<h1 style='text-align:center;margin-bottom:0'>📅 Student Timetable Generator</h1>"
    "<p style='text-align:center;margin-top:4px'>Enter your roll number to view your personal class and exam timetables.</p>",
    unsafe_allow_html=True,
)

with st.form("roll_form"):
    roll = st.text_input("Your Roll Number", placeholder="e.g. 25P060")
    st.markdown("**Compare with friends** <span style='color:#9da7b3;font-weight:400'>(optional — enter up to 5 friends)</span>", unsafe_allow_html=True)
    friend_rolls = []
    friend_cols = st.columns(3)
    for idx in range(3):
        with friend_cols[idx]:
            friend_rolls.append(st.text_input(f"Friend {idx + 1} Roll Number", placeholder="Optional", key=f"friend_roll_{idx + 1}"))
    friend_cols = st.columns(3)
    for idx in range(3, 5):
        with friend_cols[idx - 3]:
            friend_rolls.append(st.text_input(f"Friend {idx + 1} Roll Number", placeholder="Optional", key=f"friend_roll_{idx + 1}"))
    submitted = st.form_submit_button("Generate Timetables", use_container_width=True)

if submitted:
    roll = normalize_roll(roll)
    if not roll:
        st.warning("Please enter your roll number.")
        st.stop()
    try:
        records = parse_all_subject_files(DATA_DIR, tt_path)
        timetable = parse_timetable(tt_path)
        student, _, errors = build_student_timetable(records, timetable, roll)
    except Exception as e:
        st.error(f"Could not generate the timetable: {e}")
        st.stop()

    if student is None:
        st.error(f"Roll number {roll} was not found.")
        st.stop()

    students = [student]
    student["compare_label"] = "You"
    seen_rolls = {roll}
    friend_warnings = []
    for friend_roll_raw in friend_rolls:
        friend_roll = normalize_roll(friend_roll_raw)
        if not friend_roll:
            continue
        if friend_roll in seen_rolls:
            friend_warnings.append(f"Roll number {friend_roll} was entered more than once or matches your own roll number; it was skipped.")
            continue
        seen_rolls.add(friend_roll)
        friend, _, _ = build_student_timetable(records, timetable, friend_roll)
        if friend is None:
            friend_warnings.append(f"Friend roll number {friend_roll} was not found in the subject lists and was skipped.")
            continue
        friend["compare_label"] = f"Friend {len(students)}"
        students.append(friend)

    for warning in friend_warnings:
        st.warning(warning)

    exam_datesheet_path = find_exam_datesheet()
    all_exams = None
    exam_error = ""
    if exam_datesheet_path:
        try:
            all_exams = parse_exam_datesheet_pdf(exam_datesheet_path)
        except Exception as e:
            exam_error = str(e)

    tab_names = ["📚 Class Timetable", "📝 Exam Timetable"]
    if len(students) > 1:
        tab_names.append("👥 Compare Timetables")
    tabs = st.tabs(tab_names)
    class_tab, exam_tab = tabs[0], tabs[1]

    with class_tab:
        if student["missing"]:
            st.warning("Some enrolled subjects could not be matched to a class slot: " + ", ".join(student["missing"]))
        st.success(
            f'Generated class timetable for {student["name"]} ({student["roll"]}) — '
            f'{len(student["schedule"])} class entries.'
        )
        st.markdown(render_master_grid(student), unsafe_allow_html=True)

        xb, tb = make_downloads(student)
        c1, c2 = st.columns(2)
        with c1:
            st.download_button(
                "Download Excel Timetable", xb,
                file_name=f"{student['roll']}_class_timetable.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        with c2:
            st.download_button(
                "Download Copyable Timetable", tb,
                file_name=f"{student['roll']}_class_timetable.txt",
                mime="text/plain",
                use_container_width=True,
            )

    with exam_tab:
        if not exam_datesheet_path:
            st.warning(
                "The exam date-sheet PDF is not installed. Upload it to the repository's `data/` folder "
                "with a filename containing 'Exam' and 'DateSheet' (or 'Date Sheet')."
            )
        elif exam_error:
            st.error(f"Could not read the exam date sheet: {exam_error}")
        else:
            exam_rows, missing_exams, unmapped_subjects = build_student_exam_timetable(student, all_exams)
            if exam_rows:
                st.success(f"Found {len(exam_rows)} exam(s) for {student['name']} ({student['roll']}).")
            st.markdown(render_exam_table(exam_rows, student), unsafe_allow_html=True)

            if missing_exams:
                st.info("These enrolled subjects were not listed in the supplied exam date sheet: " + ", ".join(missing_exams))
            if unmapped_subjects:
                st.caption("These subject list(s) could not be matched to a known exam abbreviation: " + ", ".join(unmapped_subjects))

            if exam_rows:
                export_df = pd.DataFrame([
                    {"Day": x["day"], "Date": x["date"], "Time": x["time"], "Subject Code": x["code"], "Subject": x["subject"]}
                    for x in exam_rows
                ])
                st.download_button(
                    "Download Exam Timetable (CSV)",
                    export_df.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"{student['roll']}_exam_timetable.csv",
                    mime="text/csv",
                    use_container_width=True,
                )

    if len(students) > 1:
        compare_tab = tabs[2]
        with compare_tab:
            st.markdown("### Class schedule comparison")
            st.caption("Compare up to six students in the same day-and-period grid. Matching subjects and sections are highlighted.")
            shared_rows, common_free = comparison_summary(students)
            metric_col1, metric_col2, metric_col3 = st.columns(3)
            with metric_col1:
                st.metric("People compared", len(students))
            with metric_col2:
                st.metric("Shared class slots", len(shared_rows))
            with metric_col3:
                st.metric("Common free slots", len(common_free))

            st.markdown(render_comparison_grid(students), unsafe_allow_html=True)

            if shared_rows:
                st.markdown("#### Classes you share")
                st.dataframe(pd.DataFrame(shared_rows), use_container_width=True, hide_index=True)
            else:
                st.info("No identical subject-and-section sessions were found. The grid above still shows each person's classes and free periods.")

            if common_free:
                common_free_df = pd.DataFrame(common_free)
                with st.expander("Show all common free class slots"):
                    st.dataframe(common_free_df, use_container_width=True, hide_index=True)

            for person in students:
                if person.get("missing"):
                    who = "You" if person["roll"] == roll else f"{person.get('name') or person['roll']} ({person['roll']})"
                    st.caption(f"Subjects without a class-slot match for {who}: " + ", ".join(person["missing"]))

            st.markdown("### Exam schedule comparison")
            if exam_error:
                st.warning(f"Could not build the exam comparison: {exam_error}")
            elif not all_exams:
                st.info("An exam date-sheet PDF is not installed, so exam comparison is unavailable.")
            else:
                exam_compare_df = build_exam_comparison(students, all_exams)
                if not exam_compare_df.empty:
                    st.dataframe(exam_compare_df, use_container_width=True, hide_index=True)
                    st.download_button(
                        "Download Comparison (CSV)",
                        exam_compare_df.to_csv(index=False).encode("utf-8-sig"),
                        file_name=f"{roll}_friends_timetable_comparison.csv",
                        mime="text/csv",
                        use_container_width=True,
                    )
else:
    st.info("Enter your roll number, add up to five friends if needed, and press Enter or click Generate Timetables.")
